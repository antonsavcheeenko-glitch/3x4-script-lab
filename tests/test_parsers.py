import io

import pytest

from app.parsers import ParseError, chunk_text, parse_bytes
from app.utils.text import split_paragraphs, split_sentences


def make_pdf(text: str) -> bytes:
    """Мінімальний валідний PDF з одним рядком ASCII-тексту."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + o + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer << /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


def test_txt_utf8_and_cp1251():
    t = "Київ. 2024 рік. Тест українського тексту: ї, є, ґ, ʼ."
    assert parse_bytes("a.txt", t.encode("utf-8")).text == t
    t2 = "Привіт, це текст у кодуванні Windows-1251: їжак, єнот."
    assert parse_bytes("b.txt", t2.encode("cp1251")).text == t2


def test_markdown_strips_markup_and_takes_title():
    md = "# Заголовок\n\nТекст із **жирним** і [посиланням](https://ex.com).\n\n> Цитата"
    p = parse_bytes("x.md", md.encode())
    assert p.title == "Заголовок"
    assert "**" not in p.text and "жирним" in p.text
    assert "https://ex.com" in p.text
    assert "Цитата" in p.text and ">" not in p.text


def test_docx_parsing():
    from docx import Document

    d = Document()
    d.add_heading("Звіт про ринок", level=1)
    d.add_paragraph("У 2024 році обсяг ринку склав 12 млрд гривень.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Показник"
    t.rows[0].cells[1].text = "42"
    buf = io.BytesIO()
    d.save(buf)
    p = parse_bytes("r.docx", buf.getvalue())
    assert p.file_type == "docx"
    assert p.title == "Звіт про ринок"
    assert "12 млрд" in p.text
    assert "Показник | 42" in p.text


def test_pdf_parsing_with_pages():
    p = parse_bytes("doc.pdf", make_pdf("Casino report 2025: 62 percent"))
    assert p.file_type == "pdf"
    assert "62 percent" in p.text
    assert p.page_offsets == [(0, 1)]
    chunks = chunk_text(p.text, page_offsets=p.page_offsets)
    assert chunks[0].page == 1


def test_unsupported_and_empty():
    with pytest.raises(ParseError):
        parse_bytes("a.exe", b"xx")
    with pytest.raises(ParseError):
        parse_bytes("a.txt", b"   ")


def test_chunk_positions_match_source():
    paras = [f"Абзац номер {i}. " + "Речення з текстом. " * 20 for i in range(12)]
    text = "\n\n".join(paras)
    chunks = chunk_text(text, max_chars=900)
    assert len(chunks) > 1
    for c in chunks:
        assert text[c.start_char : c.end_char].strip() == c.text
        assert len(c.text) <= 1300
    assert [c.idx for c in chunks] == list(range(len(chunks)))


def test_sentence_splitting_handles_abbreviations():
    s = split_sentences("У 2020 р. ухвалили закон. Це було в м. Київ, т.д. Далі? Так! Кінець…")
    assert s[0] == "У 2020 р. ухвалили закон."
    assert "Далі?" in s and "Так!" in s


def test_paragraphs():
    assert split_paragraphs("a\nb\n\n\nc\n\n---\n\nd") == ["a b", "c", "d"]
