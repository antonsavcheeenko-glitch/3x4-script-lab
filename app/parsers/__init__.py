"""Парсинг файлів (.md, .txt, .docx, .pdf) у текст + розбиття на чанки."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS = {".md", ".markdown", ".txt", ".docx", ".pdf"}


class ParseError(Exception):
    pass


@dataclass
class ParsedDocument:
    text: str
    file_type: str
    title: str = ""
    # (start_char, page_number) — межі сторінок для PDF
    page_offsets: list[tuple[int, int]] = field(default_factory=list)


@dataclass
class Chunk:
    idx: int
    text: str
    start_char: int
    end_char: int
    page: int | None = None


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1251", "koi8-u"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _parse_markdown(raw: str) -> tuple[str, str]:
    title = ""
    m = re.search(r"^\s*#\s+(.+)$", raw, flags=re.M)
    if m:
        title = m.group(1).strip()
    text = raw
    text = re.sub(r"```.*?```", "", text, flags=re.S)
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)  # картинки
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)  # посилання: лишаємо URL
    text = re.sub(r"^\s{0,3}#{1,6}\s*", "", text, flags=re.M)
    text = re.sub(r"^\s*>\s?", "", text, flags=re.M)
    text = re.sub(r"(\*\*|__)(.+?)\1", r"\2", text)
    text = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip(), title


def _parse_docx(data: bytes) -> tuple[str, str]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    paras = []
    title = ""
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        if not title and p.style is not None and p.style.name.lower().startswith(("title", "heading")):
            title = t
        paras.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paras.append(" | ".join(cells))
    return "\n\n".join(paras), title


def _parse_pdf(data: bytes) -> tuple[str, list[tuple[int, int]]]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001
        raise ParseError(f"Не вдалося відкрити PDF: {e}") from e
    parts: list[str] = []
    offsets: list[tuple[int, int]] = []
    pos = 0
    for i, page in enumerate(reader.pages, start=1):
        try:
            t = page.extract_text() or ""
        except Exception:  # noqa: BLE001
            t = ""
        # PDF часто ламає рядки посеред речення: склеюємо переноси
        t = re.sub(r"-\n(?=\w)", "", t)
        t = re.sub(r"(?<![.!?:;])\n(?!\n)", " ", t)
        t = t.strip()
        offsets.append((pos, i))
        parts.append(t)
        pos += len(t) + 2
    text = "\n\n".join(parts).strip()
    if not text:
        raise ParseError("PDF не містить текстового шару (можливо, це скан). OCR у MVP не підтримується.")
    return text, offsets


def parse_bytes(filename: str, data: bytes) -> ParsedDocument:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ParseError(f"Непідтримуваний формат: {ext}. Підтримуються: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")
    title = ""
    offsets: list[tuple[int, int]] = []
    if ext in (".md", ".markdown"):
        text, title = _parse_markdown(_decode(data))
        ftype = "md"
    elif ext == ".txt":
        text = _decode(data).strip()
        ftype = "txt"
    elif ext == ".docx":
        try:
            text, title = _parse_docx(data)
        except Exception as e:  # noqa: BLE001
            raise ParseError(f"Не вдалося прочитати DOCX: {e}") from e
        ftype = "docx"
    else:
        text, offsets = _parse_pdf(data)
        ftype = "pdf"
    text = text.replace("\r\n", "\n").replace(" ", " ")
    if not text.strip():
        raise ParseError("Файл порожній або текст не знайдено.")
    return ParsedDocument(text=text, file_type=ftype, title=title or Path(filename).stem, page_offsets=offsets)


def parse_file(path: str | Path) -> ParsedDocument:
    p = Path(path)
    return parse_bytes(p.name, p.read_bytes())


def _page_for(pos: int, offsets: list[tuple[int, int]]) -> int | None:
    page = None
    for start, num in offsets:
        if start <= pos:
            page = num
        else:
            break
    return page


def chunk_text(
    text: str, max_chars: int = 1200, page_offsets: list[tuple[int, int]] | None = None
) -> list[Chunk]:
    """Ріже текст на чанки по абзацах (≈max_chars). Зберігає позиції символів у вихідному тексті."""
    spans: list[tuple[int, int]] = []
    for m in re.finditer(r"\S(?:.*?)(?=\n\s*\n|\Z)", text, flags=re.S):
        spans.append((m.start(), m.end()))
    # дуже довгі абзаци ріжемо по реченнях
    fine: list[tuple[int, int]] = []
    for s, e in spans:
        if e - s <= max_chars:
            fine.append((s, e))
            continue
        cur = s
        for m in re.finditer(r"[.!?…]+\s+", text[s:e]):
            cut = s + m.end()
            if cut - cur >= max_chars * 0.6:
                fine.append((cur, cut))
                cur = cut
        if cur < e:
            fine.append((cur, e))

    chunks: list[Chunk] = []
    group_start: int | None = None
    group_end = 0
    for s, e in fine:
        if group_start is None:
            group_start, group_end = s, e
        elif e - group_start > max_chars:
            chunks.append(_mk_chunk(len(chunks), text, group_start, group_end, page_offsets))
            group_start, group_end = s, e
        else:
            group_end = e
    if group_start is not None:
        chunks.append(_mk_chunk(len(chunks), text, group_start, group_end, page_offsets))
    return chunks


def _mk_chunk(idx: int, text: str, s: int, e: int, offsets: list[tuple[int, int]] | None) -> Chunk:
    return Chunk(idx=idx, text=text[s:e].strip(), start_char=s, end_char=e, page=_page_for(s, offsets or []))
